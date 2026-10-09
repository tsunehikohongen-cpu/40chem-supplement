"""
Build Excel with GeneSet190 CV metrics matching Table 2 format:
AUC / Accuracy / F1 / MCC (mean±SD) + TP/TN/FP/FN (mean) per fold
"""
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUTS = './outputs'   # directory holding the CV_metrics_*_GeneSet190.csv files written by run_gs3_analysis.py

# ── Load CV results ───────────────────────────────────────────────────────
nc = pd.read_csv(f'{OUTPUTS}/CV_metrics_NonCACTUS_GeneSet190.csv')
ca = pd.read_csv(f'{OUTPUTS}/CV_metrics_CACTUS_GeneSet190.csv')

def get_val(df, model, metric):
    row = df[(df['Model'] == model) & (df['Metric'] == metric)]
    if row.empty:
        return None, None
    mn = row['Mean'].values[0]
    sd = row['SD'].values[0]
    folds = [row[f'Fold{i}'].values[0] for i in range(1, 6)]
    return mn, sd, folds

# ── Styles ────────────────────────────────────────────────────────────────
FONT_NAME = 'Arial'
HDR_FILL = PatternFill('solid', start_color='1F4E78', end_color='1F4E78')
HDR_FONT = Font(name=FONT_NAME, bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=13, color='1F4E78')
SUB_FONT = Font(name=FONT_NAME, italic=True, size=10, color='595959')
BODY = Font(name=FONT_NAME, size=10)
BOLD = Font(name=FONT_NAME, bold=True, size=10)
NC_FILL = PatternFill('solid', start_color='EBF3FB', end_color='EBF3FB')   # light blue — Non-CACTUS
CA_FILL = PatternFill('solid', start_color='FEF9E7', end_color='FEF9E7')   # light yellow — CACTUS
THIN = Side(style='thin', color='BFBFBF')
BRD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT   = Alignment(horizontal='left', vertical='center')

wb = Workbook()

# ════════════════════════════════════════════════════════════════════════
# Sheet 1: Summary (Table 2 style — one model block per row-group)
# ════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = 'CV_Summary_GeneSet190'

ws['A1'] = 'GeneSet3 (Top190 genes) — Cross-validation Performance Metrics'
ws['A1'].font = TITLE_FONT
ws.merge_cells('A1:O1')
ws['A2'] = ('5-fold CV with compound-level StratifiedShuffleSplit '
            '(35 labeled compounds, test_size=0.5, random_state=42)')
ws['A2'].font = SUB_FONT
ws.merge_cells('A2:O2')

# Header row
HEADERS = [
    'Modality', 'GeneSet', 'Model',
    'AUC (mean±SD)', 'Accuracy (mean±SD)', 'F1 (mean±SD)', 'MCC (mean±SD)',
    'TP (mean)', 'TN (mean)', 'FP (mean)', 'FN (mean)',
]
hdr_r = 4
for ci, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=hdr_r, column=ci, value=h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.border = BRD; c.alignment = CENTER

MODELS = ['LR', 'SVM', 'RF']
COMBOS = [
    ('Non-CACTUS', 'GeneSet190 (190 genes)', nc, NC_FILL),
    ('CACTUS',     'GeneSet190 (200 features)', ca, CA_FILL),
]

r = hdr_r + 1
for modality, gs_label, df, fill in COMBOS:
    first_row_of_block = r
    for model in MODELS:
        def fmt(metric):
            row = df[(df['Model'] == model) & (df['Metric'] == metric)]
            if row.empty:
                return 'N/A'
            mn = row['Mean'].values[0]; sd = row['SD'].values[0]
            if metric in ('TP', 'TN', 'FP', 'FN'):
                return f'{mn:.1f}'
            return f'{mn:.3f}±{sd:.3f}'

        vals = [
            modality, gs_label, model,
            fmt('AUC'), fmt('Accuracy'), fmt('F1'), fmt('MCC'),
            fmt('TP'), fmt('TN'), fmt('FP'), fmt('FN'),
        ]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = BODY; c.fill = fill; c.border = BRD
            c.alignment = CENTER if ci > 2 else LEFT
            if ci <= 2:  # modality / gs — bold
                c.font = BOLD if ci == 1 else BODY
        r += 1

    # Merge modality and GeneSet cells across 3 model rows
    if len(MODELS) > 1:
        ws.merge_cells(start_row=first_row_of_block, start_column=1,
                       end_row=r-1, end_column=1)
        ws.merge_cells(start_row=first_row_of_block, start_column=2,
                       end_row=r-1, end_column=2)
        for col in (1, 2):
            ws.cell(row=first_row_of_block, column=col).alignment = CENTER
            ws.cell(row=first_row_of_block, column=col).font = BOLD

# Column widths
col_widths = [14, 22, 7, 17, 19, 17, 17, 10, 10, 10, 10]
for ci, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[hdr_r].height = 30
ws.freeze_panes = f'A{hdr_r+1}'

# ════════════════════════════════════════════════════════════════════════
# Sheet 2: Per-fold detail (all metrics, all 5 folds)
# ════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet('Per_Fold_Detail')
ws2['A1'] = 'Per-fold Metrics — GeneSet190'
ws2['A1'].font = TITLE_FONT
ws2.merge_cells('A1:M1')

hdr2 = ['Modality', 'Model', 'Metric',
        'Fold1', 'Fold2', 'Fold3', 'Fold4', 'Fold5', 'Mean', 'SD']
for ci, h in enumerate(hdr2, start=1):
    c = ws2.cell(row=3, column=ci, value=h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.border = BRD; c.alignment = CENTER

r2 = 4
METRICS_ORDER = ['AUC', 'Accuracy', 'F1', 'MCC', 'TP', 'TN', 'FP', 'FN']
for modality_label, gs_label, df, fill in COMBOS:
    for model in MODELS:
        for metric in METRICS_ORDER:
            row = df[(df['Model'] == model) & (df['Metric'] == metric)]
            if row.empty:
                continue
            mn = row['Mean'].values[0]; sd = row['SD'].values[0]
            folds = [row[f'Fold{i}'].values[0] for i in range(1, 6)]
            vals = [modality_label, model, metric] + folds + [mn, sd]
            for ci, v in enumerate(vals, start=1):
                c = ws2.cell(row=r2, column=ci, value=v)
                c.font = BODY; c.fill = fill; c.border = BRD
                c.alignment = CENTER if ci > 2 else LEFT
                if isinstance(v, float):
                    c.number_format = '0.000'
            r2 += 1

col_widths2 = [14, 7, 10] + [8]*5 + [8, 8]
for ci, w in enumerate(col_widths2, start=1):
    ws2.column_dimensions[get_column_letter(ci)].width = w
ws2.row_dimensions[3].height = 28
ws2.freeze_panes = 'A4'

out = f'{OUTPUTS}/GeneSet190_CV_Results.xlsx'
wb.save(out)
print(f'Saved: {out}')
