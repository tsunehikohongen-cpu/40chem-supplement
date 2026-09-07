import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

RESDIR = '/sessions/elegant-beautiful-bell/mnt/outputs/shap_results'
TAGS = ['CACTUS_GeneSet182', 'CACTUS_GeneSet150', 'NonCACTUS_GeneSet182', 'NonCACTUS_GeneSet150']
LABELS = {
    'CACTUS_GeneSet182': 'CACTUS (Multi-modal) x GeneSet182',
    'CACTUS_GeneSet150': 'CACTUS (Multi-modal) x GeneSet150',
    'NonCACTUS_GeneSet182': 'Non-CACTUS (Transcriptomics) x GeneSet182',
    'NonCACTUS_GeneSet150': 'Non-CACTUS (Transcriptomics) x GeneSet150',
}

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', start_color='1F4E78', end_color='1F4E78')
HEADER_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color='1F4E78')
SUB_FONT = Font(name=FONT, italic=True, size=10, color='595959')
BODY_FONT = Font(name=FONT, size=10)
GENE_FILL = PatternFill('solid', start_color='E8F0FE', end_color='E8F0FE')
DESC_FILL = PatternFill('solid', start_color='FCEEE3', end_color='FCEEE3')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

dfs = {tag: pd.read_csv(f'{RESDIR}/{tag}.csv') for tag in TAGS}

wb = Workbook()

# ---------------- Overview sheet ----------------
ov = wb.active
ov.title = 'Overview'
ov['A1'] = 'SHAP Feature Importance Analysis — ML Ensemble (LR / SVM / RF)'
ov['A1'].font = TITLE_FONT
ov['A2'] = 'Top contributing genes / molecular descriptors per model-modality x gene-set combination'
ov['A2'].font = SUB_FONT
ov.merge_cells('A1:F1'); ov.merge_cells('A2:F2')

ov['A4'] = 'Method'
ov['A4'].font = Font(name=FONT, bold=True, size=11)
method_lines = [
    'Final models (trained on all 35 labeled compounds, 105 concentration-level samples; StandardScaler fit on the full set) were explained with classifier-specific SHAP explainers:',
    '  - Logistic Regression (LR): shap.LinearExplainer (exact, closed-form)',
    '  - Random Forest (RF): shap.TreeExplainer (exact)',
    '  - SVM (RBF kernel): shap.KernelExplainer on predict_proba, background = 25-sample subset, nsamples = 300 (model-agnostic approximation)',
    'Importance = mean(|SHAP value|) across all 105 samples, per feature, per classifier.',
    "Ensemble_Importance = average of each classifier's mean|SHAP|, each first normalized by its own maximum (0-1 scale) so LR/RF/SVM are comparable despite differing native SHAP scales.",
]
r = 5
for line in method_lines:
    ov.cell(row=r, column=1, value=line).font = BODY_FONT
    ov.merge_cells(f'A{r}:H{r}')
    r += 1

r += 1
ov.cell(row=r, column=1, value='Combinations').font = Font(name=FONT, bold=True, size=11)
r += 1
headers = ['Combination', 'Sheet', '#Features', '#Genes', '#Descriptors', 'Top Feature (Ensemble)']
for ci, h in enumerate(headers, start=1):
    cell = ov.cell(row=r, column=ci, value=h)
    cell.font = HEADER_FONT; cell.fill = HEADER_FILL; cell.border = BORDER
    cell.alignment = Alignment(horizontal='center')
hdr_row = r
r += 1
for tag in TAGS:
    df = dfs[tag]
    n_gene = (df['FeatureType'] == 'Gene').sum()
    n_desc = (df['FeatureType'] == 'Descriptor').sum()
    top_feat = df.sort_values('Ensemble_Importance', ascending=False).iloc[0]['Feature']
    vals = [LABELS[tag], tag, len(df), int(n_gene), int(n_desc), top_feat]
    for ci, v in enumerate(vals, start=1):
        cell = ov.cell(row=r, column=ci, value=v)
        cell.font = BODY_FONT; cell.border = BORDER
    r += 1

widths = [38, 22, 11, 9, 13, 24]
for ci, w in enumerate(widths, start=1):
    ov.column_dimensions[get_column_letter(ci)].width = w
ov.freeze_panes = f'A{hdr_row+1}'

# ---------------- Detail sheets ----------------
cols = ['Rank', 'Feature', 'FeatureType', 'MeanAbsSHAP_LR', 'MeanSHAP_LR',
        'MeanAbsSHAP_RF', 'MeanSHAP_RF', 'MeanAbsSHAP_SVM', 'MeanSHAP_SVM',
        'Ensemble_Importance']
col_labels = ['Rank', 'Feature', 'Type', 'Mean|SHAP| LR', 'Mean SHAP LR',
              'Mean|SHAP| RF', 'Mean SHAP RF', 'Mean|SHAP| SVM', 'Mean SHAP SVM',
              'Ensemble Importance']

for tag in TAGS:
    df = dfs[tag][cols]
    ws = wb.create_sheet(tag[:31])
    ws['A1'] = LABELS[tag]
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:J1')
    ws['A2'] = f'{len(df)} features, ranked by Ensemble Importance (descending)'
    ws['A2'].font = SUB_FONT
    ws.merge_cells('A2:J2')

    header_row = 4
    for ci, h in enumerate(col_labels, start=1):
        cell = ws.cell(row=header_row, column=ci, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL; cell.border = BORDER
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    for ri, row in enumerate(df.itertuples(index=False), start=header_row+1):
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            if ci in (4, 5, 6, 7, 8, 9, 10) and isinstance(val, float):
                cell.number_format = '0.000000'
            if ci == 3:
                cell.fill = GENE_FILL if val == 'Gene' else DESC_FILL
            cell.alignment = Alignment(horizontal='center' if ci != 2 else 'left')

    last_row = header_row + len(df)
    rule = ColorScaleRule(start_type='min', start_color='FFFFFF',
                           end_type='max', end_color='63BE7B')
    ws.conditional_formatting.add(f'J{header_row+1}:J{last_row}', rule)

    widths = [6, 16, 11, 13, 12, 13, 12, 13, 12, 16]
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = f'A{header_row+1}'
    ws.auto_filter.ref = f'A{header_row}:J{last_row}'

out_path = '/sessions/elegant-beautiful-bell/mnt/outputs/SHAP_Feature_Importance.xlsx'
wb.save(out_path)
print('saved', out_path)
